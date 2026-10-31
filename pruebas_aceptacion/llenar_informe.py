"""Completa el documento 4.2 (Pruebas de aceptación) con las evidencias de Selenium.

Uso:
    python llenar_informe.py 4_2_Pruebas_de_Aceptacion.docx
    python llenar_informe.py 4_2_Pruebas_de_Aceptacion.docx --evidencias evidencias --salida Informe_final.docx

Qué hace, para cada «Prueba CAxxx» de la sección 4.2.2:
  * Imagen de la prueba  -> los pantallazos numerados paso a paso (evidencias/CAxxx/NN_*.png)
                            + cualquier evidencias/CAxxx/manual_*.png que usted agregue a mano.
  * Resultado            -> resultado obtenido y pasos verificados.
  * Observaciones        -> hallazgos técnicos/funcionales detectados durante la ejecución.
  * Concepto Final       -> Aprobado / Aprobado con observaciones / No aprobado.
También reemplaza «[Pendiente de ejecución; ver 4.2.2]» en las tablas de diseño (4.2.1).
Lee evidencias/resultados.json, que genera `pytest`.
"""
import argparse
import datetime
import json
import pathlib
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ANCHO_IMG = Cm(8.6)
COL_CM = (3.0, 9.0, 4.5)  # caso | imagen | resultado  (suma 16.5 cm = ancho útil)


def concepto(caso):
    r = caso["resultado"]
    if r == "No aprobado":
        return "No aprobado"
    if r in ("Omitido", "Parcial"):
        return "Pendiente (ejecución incompleta)"
    return "Aprobado con observaciones" if caso["observaciones"] else "Aprobado"


def limpiar(celda):
    """Deja la celda con un único párrafo vacío (conserva el formato de la celda)."""
    for p in celda.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = celda.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    return p


def escribir(celda, lineas, tam=8, negrita_primera=False, color=None):
    """Escribe una lista de líneas, una por párrafo."""
    p = limpiar(celda)
    for i, linea in enumerate(lineas):
        if i:
            p = celda.add_paragraph()
        run = p.add_run(linea)
        run.font.size = Pt(tam)
        run.bold = bool(negrita_primera and i == 0)
        if color:
            run.font.color.rgb = color
        p.paragraph_format.space_after = Pt(2)


def poner_imagenes(celda, items):
    """items: [(leyenda, ruta_png)]"""
    p = limpiar(celda)
    primera = True
    for leyenda, ruta in items:
        if not primera:
            p = celda.add_paragraph()
        primera = False
        run = p.add_run(leyenda)
        run.font.size = Pt(7.5)
        run.italic = True
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        img = celda.add_paragraph()
        img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img.paragraph_format.space_after = Pt(6)
        img.add_run().add_picture(str(ruta), width=ANCHO_IMG)
        p = img


def fijar_anchos(tabla):
    for col, cm in zip(tabla.columns, COL_CM):
        col.width = Cm(cm)
    for fila in tabla.rows:
        tcs = []
        for c in fila.cells:
            if not any(c._tc is t for t in tcs):
                tcs.append(c._tc)
        celdas = {id(c._tc): c for c in fila.cells}.values()
        celdas = list(celdas)
        if len(celdas) == 3:
            for c, cm in zip(celdas, COL_CM):
                c.width = Cm(cm)
        elif len(celdas) == 2:
            celdas[0].width = Cm(COL_CM[0])
            celdas[1].width = Cm(COL_CM[1] + COL_CM[2])


def imagenes_del_caso(caso):
    items = []
    for paso in caso["pasos"]:
        if pathlib.Path(paso["archivo"]).exists():
            items.append((f"Paso {paso['n']}. {paso['texto']}", paso["archivo"]))
    if caso["resultado"] == "No aprobado" or not items:
        for ruta in caso["captura_final"]:
            if pathlib.Path(ruta).exists():
                etiqueta = "Pantallazo del fallo" if "FALLO" in ruta else "Pantallazo final"
                items.append((etiqueta, ruta))
    for ruta in caso["manuales"]:
        if pathlib.Path(ruta).exists():
            items.append((f"Evidencia adicional ({pathlib.Path(ruta).name})", ruta))
    return items


def texto_resultado(caso, fecha):
    r = caso["resultado"]
    if r == "Aprobado":
        cab = "Exitoso. La ejecución automatizada con Selenium cumplió el resultado esperado."
    elif r == "No aprobado":
        cab = "Fallido. La ejecución automatizada no cumplió el resultado esperado."
    else:
        cab = "No se pudo completar la ejecución automatizada."
    lineas = [cab, f"Fecha de ejecución: {fecha}."]
    if caso["pasos"]:
        lineas.append("Pasos verificados:")
        lineas += [f"{p['n']}. {p['texto']}" for p in caso["pasos"]]
    if r != "Aprobado" and caso["detalle"]:
        lineas.append(f"Detalle: {caso['detalle']}")
    return lineas


BASE = pathlib.Path(__file__).resolve().parent.parent  # raíz del proyecto


def ruta_segura(valor):
    """Resuelve una ruta recibida por la línea de comandos y exige que esté dentro del proyecto."""
    ruta = (pathlib.Path.cwd() / valor).resolve()
    if not ruta.is_relative_to(BASE):
        sys.exit(f"Ruta fuera del proyecto: {valor}")
    return ruta


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("docx", help="Documento 4.2 original")
    ap.add_argument("--evidencias", default="evidencias")
    ap.add_argument("--salida", default=None)
    a = ap.parse_args()
    a.docx = str(ruta_segura(a.docx))
    a.evidencias = str(ruta_segura(a.evidencias))
    if a.salida:
        a.salida = str(ruta_segura(a.salida))

    ruta_json = pathlib.Path(a.evidencias) / "resultados.json"
    if not ruta_json.exists():
        sys.exit(f"No existe {ruta_json}. Ejecute primero: pytest")
    datos = json.loads(ruta_json.read_text(encoding="utf-8"))
    casos = datos["casos"]
    fecha = datetime.datetime.fromisoformat(datos["generado"]).strftime("%d/%m/%Y")

    doc = Document(a.docx)
    llenos, pendientes = [], []
    for tabla in doc.tables:
        titulo = tabla.rows[0].cells[0].text.strip()
        m = re.fullmatch(r"Prueba (CA\d{3})", titulo)
        if m:  # ── 4.2.2 Ejecución y evaluación
            id_caso = m.group(1)
            caso = casos.get(id_caso)
            if not caso:
                pendientes.append(id_caso)
                continue
            fijar_anchos(tabla)
            items = imagenes_del_caso(caso)
            if items:
                poner_imagenes(tabla.rows[2].cells[1], items)
            else:
                escribir(tabla.rows[2].cells[1], ["(sin capturas: la prueba no llegó a ejecutarse)"])
            escribir(tabla.rows[2].cells[2], texto_resultado(caso, fecha))
            if caso["observaciones"]:
                obs = caso["observaciones"]
            elif caso["resultado"] == "Aprobado":
                obs = ["Sin diferencias frente al resultado esperado."]
            else:
                obs = [caso["detalle"] or "Prueba no ejecutada."]
            escribir(tabla.rows[3].cells[1], obs)
            con = concepto(caso)
            color = {"Aprobado": RGBColor(0x0B, 0x7A, 0x3B),
                     "Aprobado con observaciones": RGBColor(0xB4, 0x6A, 0x00),
                     "No aprobado": RGBColor(0xB0, 0x1E, 0x1E)}.get(con)
            escribir(tabla.rows[4].cells[1], [con], tam=9, negrita_primera=True, color=color)
            (pendientes if caso["resultado"] in ("Omitido", "Parcial") else llenos).append(id_caso)
            continue

        # ── 4.2.1 Diseño: «Resultado Obtenido»
        m = re.fullmatch(r"ID Caso:\s*(CA\d{3})", titulo)
        if m and m.group(1) in casos:
            caso = casos[m.group(1)]
            frase = {"Aprobado": "Según lo esperado; ver 4.2.2",
                     "No aprobado": "No coincide con lo esperado; ver 4.2.2"}.get(
                caso["resultado"], "Ejecución incompleta; ver 4.2.2")
            if caso["resultado"] == "Aprobado" and caso["observaciones"]:
                frase = "Según lo esperado, con observaciones; ver 4.2.2"
            for fila in tabla.rows:
                for c in {id(x._tc): x for x in fila.cells}.values():
                    for p in c.paragraphs:
                        for r in p.runs:
                            if "[Pendiente de ejecución; ver 4.2.2]" in r.text:
                                r.text = r.text.replace("[Pendiente de ejecución; ver 4.2.2]", frase)

    salida = a.salida or str(pathlib.Path(a.docx).with_name(
        pathlib.Path(a.docx).stem + "_con_evidencias.docx"))
    doc.save(salida)
    print(f"Documento generado: {salida}")
    print(f"  Casos completados : {len(llenos)} {sorted(llenos)}")
    if pendientes:
        print(f"  Sin completar     : {len(pendientes)} {sorted(pendientes)}  (omitidos, parciales o sin ejecutar)")


if __name__ == "__main__":
    main()
